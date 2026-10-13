# A superscript footnote marker run ("5%" + "1" + " in 2023") used to be
# re-flowed by length: the "1" lost its superscript and the space got it.
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches

from core.pipelines.ppt_translation_pipeline import (
    extract_ppt_content_to_json, write_translated_content_to_ppt)


def test_superscript_marker_stays_on_its_run(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for text, sup in [("Revenue grew 5%", False), ("1", True), (" in 2023 across all regions", False)]:
        r = p.add_run()
        r.text = text
        if sup:
            r.font._rPr.set("baseline", "30000")
    src = str(tmp_path / "in.pptx")
    prs.save(src)

    src_json = extract_ppt_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    assert [i["value"] for i in items] == ["Revenue grew 5%{{SUP_1}} in 2023 across all regions"]
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": items[0]["count_src"],
                    "translated": "收入增长了5%{{SUP_1}}，覆盖2023年所有地区"}], f, ensure_ascii=False)
    result = write_translated_content_to_ppt(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        src_lang="en", dst_lang="zh")

    runs = Presentation(result).slides[0].shapes[0].text_frame.paragraphs[0].runs
    assert [(r.text, r.font._rPr.get("baseline")) for r in runs if r.text] == [
        ("收入增长了5%", None), ("1", "30000"), ("，覆盖2023年所有地区", None)]
