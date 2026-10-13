# Word inline structures that used to drift or break on write-back:
#  - text inside w:ins / w:smartTag / inline w:sdt was never extracted, stayed
#    untranslated and jumped to the paragraph start; w:del text did the same;
#  - bookmarks collapsed to an empty range in front of the rebuilt runs, so a
#    REF cross-reference updated to "";
#  - NUMPAGES got the {{PAGE_NUMBER}} placeholder (substring match), so a
#    translation reordering "page X of Y" swapped the two fields.
# Deterministic: "translation" = uppercase outside {{...}} placeholders.
import json
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from core.pipelines.word_translation_pipeline import (
    extract_word_content_to_json, write_translated_content_to_word)

PARAS = [
    '<w:r><w:t xml:space="preserve">Hello </w:t></w:r>'
    '<w:ins w:id="1" w:author="a"><w:r><w:t>inserted world</w:t></w:r></w:ins>'
    '<w:del w:id="2" w:author="a"><w:r><w:delText>gone</w:delText></w:r></w:del>'
    '<w:r><w:t xml:space="preserve"> end text</w:t></w:r>',
    '<w:r><w:t xml:space="preserve">Meet in </w:t></w:r>'
    '<w:smartTag w:uri="u" w:element="City"><w:r><w:t>Paris</w:t></w:r></w:smartTag>'
    '<w:r><w:t xml:space="preserve"> tomorrow please</w:t></w:r>',
    '<w:r><w:t xml:space="preserve">Name: </w:t></w:r>'
    '<w:sdt><w:sdtPr><w:id w:val="5"/></w:sdtPr><w:sdtContent>'
    '<w:r><w:t>Click here to enter</w:t></w:r></w:sdtContent></w:sdt>'
    '<w:r><w:t xml:space="preserve"> and sign below</w:t></w:r>',
    '<w:bookmarkStart w:id="0" w:name="_Ref1"/><w:r><w:t>Installation guide overview</w:t></w:r>'
    '<w:bookmarkEnd w:id="0"/>',
    '<w:r><w:t xml:space="preserve">Company confidential page </w:t></w:r>'
    '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple>'
    '<w:r><w:t xml:space="preserve"> of </w:t></w:r>'
    '<w:fldSimple w:instr=" NUMPAGES "><w:r><w:t>5</w:t></w:r></w:fldSimple>',
]


def _fake(text):
    if "{{TOTAL_PAGES}}" in text:          # target language puts the total first
        return "TOTAL {{TOTAL_PAGES}} - PAGE {{PAGE_NUMBER}}"
    return "".join(p if p.startswith("{{") else p.upper()
                   for p in re.split(r"(\{\{.*?\}\})", text))


def _roundtrip(tmp_path):
    src = str(tmp_path / "inline.docx")
    doc = Document()
    body = doc.element.body
    for xml in PARAS:
        body.insert(len(body) - 1, parse_xml(f"<w:p {nsdecls('w')}>{xml}</w:p>"))
    doc.save(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": i["count_src"], "translated": _fake(i["value"])} for i in items],
                  f, ensure_ascii=False)
    result = write_translated_content_to_word(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        src_lang="en", dst_lang="es")
    return [i["value"] for i in items], Document(result).paragraphs


def _text(p):
    return "".join(t.text or "" for t in p._p.iter(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"))


def test_wrapped_text_translated_in_place(tmp_path):
    values, paras = _roundtrip(tmp_path)
    assert "Hello {{HLINK_1}}inserted world{{/HLINK_1}} end text" in values
    assert [_text(p) for p in paras[:3]] == [
        "HELLO INSERTED WORLD END TEXT", "MEET IN PARIS TOMORROW PLEASE",
        "NAME: CLICK HERE TO ENTER AND SIGN BELOW"]
    xml = [p._p.xml for p in paras[:3]]
    assert "<w:ins" in xml[0] and "w:del" not in xml[0] and "gone" not in xml[0]
    assert "<w:smartTag" in xml[1] and "<w:sdtContent>" in xml[2]
    assert re.search(r"<w:sdtContent>.*CLICK HERE TO ENTER.*</w:sdtContent>", xml[2], re.S)
    assert all("{{" not in x for x in xml)


def test_bookmark_spans_translated_text(tmp_path):
    _, paras = _roundtrip(tmp_path)
    xml = paras[3]._p.xml
    assert re.search(r"<w:bookmarkStart[^>]*_Ref1.*INSTALLATION GUIDE OVERVIEW.*<w:bookmarkEnd", xml, re.S)


def test_numpages_has_own_placeholder(tmp_path):
    values, paras = _roundtrip(tmp_path)
    assert "Company confidential page {{PAGE_NUMBER}} of {{TOTAL_PAGES}}" in values
    instrs = re.findall(r'w:instr="([^"]+)"', paras[4]._p.xml)
    assert [i.strip() for i in instrs] == ["NUMPAGES", "PAGE"]
