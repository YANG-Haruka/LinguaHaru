# Shared engine: a batch is a JSON {id: text}. A model can return every id but
# with content shifted by one (id1 = text1+text2, id2 = text3, ...); each item
# alone looks valid, so the whole document drifted silently. These tests drive
# the REAL txt / docx translators with a fake model (no network).
#
#   python -m pytest tests/test_engine_batch_alignment.py
import json
import os
import sys

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, REPO_ROOT)
os.chdir(REPO_ROOT)

import pytest

import core.engine.base_translator as bt
from core.engine.translation_checker import _batch_misaligned, clean_json
from core.engine import translation_qa as qa

LINES = ["The meeting starts at 9 am.", "Bring your laptop and charger.",
         "Lunch will be served at noon.", "The budget is 4500 dollars.",
         "Questions go to Alice.", "Thanks for coming."]


def zh(t):
    return "译:" + t


def merge(vals):      # id1 = t1+t2, later ids one early, last slot repeated
    return [vals[0] + " " + vals[1]] + vals[2:] + [vals[-1]]


def shift(vals):      # every id one early, last slot empty
    return vals[1:] + [""]


def _run(tmp_path, monkeypatch, drift, ext):
    def fake(segments, *a, **k):
        d = json.loads(clean_json(segments))
        keys, vals = list(d), list(d.values())
        out = vals if len(keys) == 1 else drift(vals)
        return json.dumps({k: zh(v) if v else "" for k, v in zip(keys, out)},
                          ensure_ascii=False), True, None
    monkeypatch.setattr(bt, "translate_text", fake)
    monkeypatch.setattr(bt, "cache_store_validated", lambda *a, **k: None)
    src = str(tmp_path / f"a{ext}")
    if ext == ".txt":
        from core.translators.txt_translator import TxtTranslator as cls
        with open(src, "w", encoding="utf-8") as f:
            f.write("\n".join(LINES))
    else:
        import docx
        from core.translators.word_translator import WordTranslator as cls
        doc = docx.Document()
        for line in LINES:
            doc.add_paragraph(line)
        doc.save(src)
    d = {n: str(tmp_path / n) for n in ("temp", "result", "log", "history")}
    for p in d.values():
        os.makedirs(p, exist_ok=True)
    tr = cls(src, "fake", False, "", "en", "zh", False, max_token=768, max_retries=3,
             thread_count=1, glossary_path=None, temp_dir=d["temp"], result_dir=d["result"],
             session_lang="en", log_dir=d["log"], history_dir=d["history"])
    out, _ = tr.process(os.path.splitext(src)[0], ext)
    if ext == ".txt":
        with open(out, encoding="utf-8-sig") as f:
            return [x for x in f.read().splitlines() if x.strip()]
    import docx
    return [p.text for p in docx.Document(out).paragraphs if p.text.strip()]


@pytest.mark.parametrize("ext", [".txt", ".docx"])
@pytest.mark.parametrize("drift", [merge, shift])
def test_shifted_batch_is_retried_not_written(tmp_path, monkeypatch, drift, ext):
    assert _run(tmp_path, monkeypatch, drift, ext) == [zh(x) for x in LINES]


def test_guard_signals():
    src = LINES
    assert _batch_misaligned(list(zip(src, map(zh, merge(src)))))
    assert _batch_misaligned(list(zip(src, map(zh, shift(src)))))
    # correct batch, incl. localized numbers and a short repeated answer
    ok = ["会议上午9点开始。", "带上笔记本和充电器。", "中午供应午餐。",
          "预算为4,500美元。", "问题请找爱丽丝。", "感谢光临。"]
    assert not _batch_misaligned(list(zip(src, ok)))
    assert not _batch_misaligned([("Yes", "是"), ("Yeah", "是")])
    # varied lengths: length-only signal (no numbers, no duplicates)
    s = ["Hi there.", "This paragraph is quite a bit longer than the first one is.",
         "Short line.", "Another fairly long sentence that keeps going for a while here.",
         "Okay then.", "Medium length sentence here."]
    good = [x.upper() for x in s]
    assert not _batch_misaligned(list(zip(s, good)))
    assert _batch_misaligned(list(zip(s, good[1:] + [good[-1] + "!"])))


def test_length_ratio_counts_cjk_width():
    pairs = [(1, "The International Business Machines Corporation announced "
                 "quarterly earnings.", "国际商业机器公司公布了季度收益。", "text"),
             (2, "我们明天去北京参观故宫博物院。",
                 "Tomorrow we are all going to Beijing to visit the Palace Museum.", "text")]
    assert qa.check_length_ratio(pairs) == []
    assert qa.check_length_ratio([(3, "A long English sentence here.", "a", "text")]) == [3]


def test_glossary_qa_matches_like_the_prompt():
    g = [("app", "应用程序"), ("count", "计数")]
    assert qa.check_glossary_terms([(1, "Open the application settings.", "打开应用设置。", "t")], g) == []
    assert qa.check_glossary_terms([(1, "You have {count} items.", "您有{count}件商品。", "t")], g) == []
    assert qa.check_glossary_terms([(1, "Open the app.", "打开软件。", "t")], g) == \
        [{"id": 1, "term": "app", "expected": "应用程序"}]
