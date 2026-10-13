# Excel write-back bugs:
#  - two long translated sheet names truncated to the same 31 chars: openpyxl's
#    de-dup appended "1" -> 32-char name -> Excel refuses to open the workbook;
#  - rich-text cells (CellRichText runs) were flattened on save, even cells that
#    were never translated ("12" + superscript "3" became plain "123").
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from core.pipelines.excel_translation_pipeline import (
    extract_excel_content_to_json, write_translated_content_to_excel)

LONG = "很长的名字" * 8
TR = {"Report One": LONG, "Report Two": LONG + "二", "Overdue amount": "逾期金额"}


def _build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report One"
    ws["A1"] = CellRichText([TextBlock(InlineFont(b=True), "12"),
                             TextBlock(InlineFont(vertAlign="superscript"), "3")])
    ws["A2"] = CellRichText([TextBlock(InlineFont(b=True, color="FFFF0000"), "Overdue amount")])
    wb.create_sheet("Report Two")["A1"] = "='Report One'!A1"
    wb.save(path)


def _translate(tmp_path):
    src = str(tmp_path / "book.xlsx")
    _build(src)
    src_json = extract_excel_content_to_json(src, str(tmp_path / "tmp"), use_xlwings=False)
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": i["count_src"], "translated": TR.get(i["value"], i["value"])}
                   for i in items], f, ensure_ascii=False)
    out = write_translated_content_to_excel(src, src_json, dst_json, str(tmp_path / "out"),
                                            src_lang="en", dst_lang="zh")
    return load_workbook(out, rich_text=True)


def test_truncated_sheet_names_unique_within_31(tmp_path):
    wb = _translate(tmp_path)
    names = wb.sheetnames
    assert all(len(n) <= 31 for n in names), names
    assert len({n.lower() for n in names}) == 2
    assert wb[names[1]]["A1"].value == f"='{names[0]}'!A1"


def test_rich_text_runs_kept(tmp_path):
    ws = _translate(tmp_path).worksheets[0]
    a1 = ws["A1"].value
    assert isinstance(a1, CellRichText)
    assert [(b.text, b.font.b, b.font.vertAlign) for b in a1] == [("12", True, None), ("3", False, "superscript")]
    a2 = ws["A2"].value              # single-font rich cell keeps its font when translated
    assert isinstance(a2, CellRichText) and str(a2) == "逾期金额"
    assert a2[0].font.b and a2[0].font.color.rgb == "FFFF0000"
