# A multi-letter superscript ("TM", "th") and a hyperlink run used to be
# re-flowed by length, landing on whatever translated text fell at their old
# offset ("BrandX™ is..." -> "Bra" + superscript "nd"; link on "访问我们的网").
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches

from core.pipelines.ppt_translation_pipeline import (
    extract_ppt_content_to_json, write_translated_content_to_ppt)


def _roundtrip(tmp_path, runs, translated):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    p = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.paragraphs[0]
    for text, kind in runs:
        r = p.add_run()
        r.text = text
        if kind == "sup":
            r.font._rPr.set("baseline", "30000")
        elif kind == "link":
            r.hyperlink.address = "https://example.com"
    src = str(tmp_path / "in.pptx")
    prs.save(src)
    src_json = extract_ppt_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": items[0]["count_src"], "translated": translated}], f, ensure_ascii=False)
    result = write_translated_content_to_ppt(src, src_json, dst_json, str(tmp_path / "temp"),
                                             str(tmp_path / "result"))
    return items[0]["value"], Presentation(result).slides[0].shapes[0].text_frame.paragraphs[0].runs


def test_multi_letter_superscript_not_moved(tmp_path):
    value, runs = _roundtrip(tmp_path, [("BrandX", None), ("TM", "sup"),
                                        (" is a registered product name", None)],
                             "BrandX是注册产品名称")
    assert value == "BrandXTM is a registered product name"
    assert "".join(r.text for r in runs) == "BrandX是注册产品名称"
    assert all(not r.text for r in runs if r.font._rPr.get("baseline"))


def test_hyperlink_stays_on_link_text(tmp_path):
    value, runs = _roundtrip(tmp_path, [("Visit ", None), ("our website", "link"),
                                        (" for more details today", None)],
                             "今天请访问{{HLINK_1}}我们的网站{{/HLINK_1}}了解更多详情")
    assert value == "Visit {{HLINK_1}}our website{{/HLINK_1}} for more details today"
    assert [(r.text, bool(r.hyperlink.address)) for r in runs] == [
        ("今天请访问", False), ("我们的网站", True), ("了解更多详情", False)]


def test_dropped_link_marker_falls_back_without_marker_text(tmp_path):
    _, runs = _roundtrip(tmp_path, [("Visit ", None), ("our website", "link"),
                                    (" for more details today", None)],
                         "今天请访问我们的网站了解更多详情")
    assert "".join(r.text for r in runs) == "今天请访问我们的网站了解更多详情"
