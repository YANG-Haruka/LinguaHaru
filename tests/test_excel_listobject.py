# Excel Tables (ListObjects): <tableColumn name> must equal the translated header
# cell, or Excel "repairs" the workbook and reverts the headers.
import json
import os
import re
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableFormula

from core.pipelines.excel_translation_pipeline import (
    _rename_structured_refs, extract_excel_content_to_json, write_translated_content_to_excel)

HEADERS = ["Year", "Region", "Amount", "Total", "Line one\nline two"]
TR = {"Year": "年份", "Region": "地区", "Amount": "金额", "Total": "金额",   # Total collides -> 金额2
      "Sum": "合计", "North": "北方", "South": "南方", "Line one\nline two": "第一行\n第二行"}


def _build(path):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for i, region in enumerate(["North", "South", "North"]):
        ws.append([2020 + i, region, 10 * (i + 1), "=Sales[[#This Row],[Amount]]*2", 1])
    ws.append(["Sum", None, "=SUBTOTAL(109,Sales[Amount])", None, None])
    ws["H1"] = "=SUM(Sales[Amount])+SUM(Sales[[#All],[Line one\nline two]])"
    t = Table(displayName="Sales", ref="A1:E5", totalsRowCount=1)
    t._initialise_columns()
    for col, h in zip(t.tableColumns, HEADERS):
        col.name = h.replace("\n", "_x000a_")          # as Excel stores it
    t.tableColumns[0].totalsRowLabel = "Sum"
    t.tableColumns[2].totalsRowFunction = "sum"
    t.tableColumns[3].calculatedColumnFormula = TableFormula(attr_text="Sales[[#This Row],[Amount]]*2")
    t.autoFilter.ref = "A1:E4"
    ws.add_table(t)
    wb.save(path)


def _translate(src, out_dir, bilingual):
    """Real extract + write-back; a fixed map stands in for the LLM."""
    tmp = os.path.join(out_dir, "tmp")
    os.makedirs(tmp, exist_ok=True)
    src_json = extract_excel_content_to_json(src, tmp, use_xlwings=False)
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst = [{"count_src": it["count_src"],
            "translated": TR.get(it["value"].replace("␊", "\n"), it["value"]).replace("\n", "␊")}
           for it in items]
    dst_json = os.path.join(tmp, "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump(dst, f, ensure_ascii=False)
    return write_translated_content_to_excel(src, src_json, dst_json, out_dir,
                                             bilingual_mode=bilingual, src_lang="en", dst_lang="zh")


@pytest.mark.parametrize("bilingual", [False, True])
def test_table_columns_follow_translated_headers(tmp_path, bilingual):
    src = str(tmp_path / "table.xlsx")
    _build(src)
    out = _translate(src, str(tmp_path / ("bi" if bilingual else "mono")), bilingual)

    wb = load_workbook(out)                      # package reopens cleanly
    ws = wb.active
    t = ws.tables["Sales"]
    assert t.displayName == "Sales"              # identifier untouched
    headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
    names = [c.name for c in t.tableColumns]     # openpyxl decodes _x000a_
    assert names == headers
    assert len({n.lower() for n in names}) == len(names)

    if bilingual:
        assert headers[0] == "Year\n年份"
        amount = "Amount\n金额"
    else:
        assert headers == ["年份", "地区", "金额", "金额2", "第一行\n第二行"]
        amount = "金额"
    assert ws["H1"].value == f"=SUM(Sales[{amount}])+SUM(Sales[[#All],[{headers[4]}]])"
    assert ws["C5"].value == f"=SUBTOTAL(109,Sales[{amount}])"
    assert ws["D2"].value == f"=Sales[[#This Row],[{amount}]]*2"
    assert t.tableColumns[3].calculatedColumnFormula.attr_text == f"Sales[[#This Row],[{amount}]]*2"

    with zipfile.ZipFile(out) as z:
        xml = z.read("xl/tables/table1.xml").decode("utf-8")
    assert 'displayName="Sales"' in xml and 'ref="A1:E5"' in xml
    esc = [h.replace("\n", "_x000a_") for h in headers]          # Excel form: LF -> _x000a_
    assert re.findall(r'<tableColumn [^>]*name="([^"]*)"', xml) == esc
    assert re.findall(r'totalsRowLabel="([^"]*)"', xml) == [ws["A5"].value.replace("\n", "_x000a_")]

    _excel_check(out, headers)


def test_structured_ref_escaping():
    r = {"t": {"q[1]#'x": "a]b", "amount": "金额"}}
    assert _rename_structured_refs("SUM(T[[#Data],[Q'[1']'#''x]])", r) == "SUM(T[[#Data],[a']b]])"
    assert _rename_structured_refs("T[Amount]+Other[Amount]+T[#All]", r) == "T[金额]+Other[Amount]+T[#All]"


def test_external_workbook_ref_untouched():
    r = {"sales": {"amount": "金额"}}
    assert _rename_structured_refs("[1]!Sales[Amount]+Sales[Amount]", r) == "[1]!Sales[Amount]+Sales[金额]"


def test_validation_cf_and_pivot_follow_renamed_columns():
    from types import SimpleNamespace as NS
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from core.pipelines.excel_translation_pipeline import _sync_table_columns

    wb = Workbook()
    ws = wb.active
    ws.append(["地区", "金额"])                       # headers already translated
    ws.append(["北方", 1])
    t = Table(displayName="Sales", ref="A1:B2")
    t._initialise_columns()
    t.tableColumns[0].name, t.tableColumns[1].name = "Region", "Amount"
    ws.add_table(t)
    dv = DataValidation(type="list", formula1='INDIRECT("Sales[Region]")')
    ws.add_data_validation(dv)
    ws.conditional_formatting.add("D1", FormulaRule(formula=['D1>SUM(INDIRECT("Sales[Amount]"))']))
    fields = [NS(name="Region"), NS(name="Amount")]   # stands in for an openpyxl pivot cache
    ws._pivots.append(NS(cache=NS(cacheSource=NS(worksheetSource=NS(name="Sales")),
                                  cacheFields=fields)))

    _sync_table_columns(wb)

    assert dv.formula1 == 'INDIRECT("Sales[地区]")'
    rule = next(iter(ws.conditional_formatting)).rules[0]
    assert rule.formula == ['D1>SUM(INDIRECT("Sales[金额]"))']
    assert [f.name for f in fields] == ["地区", "金额"]


def _excel_check(path, headers):
    """Optional: open in real Excel (Windows + COM); a repair reverts the
    headers / column names, so equality here means Excel accepted the file."""
    try:
        import win32com.client
        xl = win32com.client.DispatchEx("Excel.Application")
    except Exception:  # noqa: BLE001
        return
    try:
        xl.DisplayAlerts = False
        wb = xl.Workbooks.Open(os.path.abspath(path))
        try:
            ws = wb.Worksheets(1)
            lo = ws.ListObjects("Sales")
            assert [lo.ListColumns(i).Name for i in range(1, 6)] == headers
            assert [ws.Cells(1, i).Value for i in range(1, 6)] == headers
            assert "repair" not in wb.Name.lower() and "修复" not in wb.Name
        finally:
            wb.Close(False)
    finally:
        xl.Quit()
