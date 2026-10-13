# Real-document bugs from the USGS "Lithium 2024" fact sheet (Word):
#  - superscript footnote markers mixed with normal runs got re-flowed by length
#    ("ᵉ2,630" -> "²630", "Consumo, ᵃᵖᵃʳᵉⁿᵗᵉ¹");
#  - a data cell holding just "W" (withheld) was translated to "西"/"O".
# Deterministic: a fake translation map stands in for the LLM.
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from core.pipelines.word_translation_pipeline import (
    extract_word_content_to_json, write_translated_content_to_word)


def _runs(cell_runs, para):
    for text, sup, bold in cell_runs:
        r = para.add_run(text)
        r.font.superscript = sup or None
        r.bold = bold or None


def _build(path):
    doc = Document()
    p = doc.add_paragraph()
    _runs([("World Resources:", False, True), ("4", True, False),
           (" Owing to continuing exploration.", False, False)], p)
    cells = [
        [("e", True, False), ("2,630", False, False)],
        [("W", False, False)],
        [("7", True, False), ("146,000", False, False)],
        [("Consumption, apparent", False, False), ("1", True, False)],
        [("Net import reliance", False, False), ("3", True, False),
         (" as a percentage of apparent consumption", False, False)],
        [("2023", False, False), ("e", True, False)],
        [("NA", False, False)],
        [("—", False, False)],
        [("Reserves", False, False)],
    ]
    table = doc.add_table(rows=1, cols=len(cells))
    for cell, spec in zip(table.rows[0].cells, cells):
        _runs(spec, cell.paragraphs[0])

    # Hyperlink + marker: exercises the generic field path
    hp = doc.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), "x")
    lr = OxmlElement("w:r")
    lt = OxmlElement("w:t")
    lt.text = "link"
    lr.append(lt)
    link.append(lr)
    hp._p.append(link)
    _runs([("Text", False, False), ("5", True, False)], hp)
    doc.save(path)


TRANSLATIONS = {
    "World Resources:{{SUP_1}} Owing to continuing exploration.":
        "世界资源：{{SUP_1}} 由于持续勘探。",
    "Consumption, apparent{{SUP_1}}": "Consumo, aparente{{SUP_1}}",
    "Net import reliance{{SUP_1}} as a percentage of apparent consumption":
        "Dependencia neta de importaciones{{SUP_1}} como porcentaje del consumo aparente",
    "Reserves": "Reservas",
    "{{HLINK_1}}link{{/HLINK_1}}Text{{SUP_1}}": "{{HLINK_1}}enlace{{/HLINK_1}}Texto{{SUP_1}}",
}


def _roundtrip(tmp_path, bilingual=False):
    src = str(tmp_path / "lithium.docx")
    _build(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    values = [i["value"] for i in items]
    out = [{"count_src": i["count_src"], "translated": TRANSLATIONS[i["value"]]} for i in items]
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False)
    result = write_translated_content_to_word(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        bilingual_mode=bilingual, src_lang="en", dst_lang="es")
    return values, Document(result)


def _sig(para):
    return [(r.text, bool(r.font.superscript), bool(r.bold)) for r in para.runs if r.text]


def test_short_and_numeric_cells_not_sent(tmp_path):
    values, _ = _roundtrip(tmp_path)
    # Exactly the translatable items; "W", "NA", "—", numbers and
    # number+marker cells never reach the translator.
    assert sorted(values) == sorted(TRANSLATIONS)


def test_superscript_markers_stay_on_their_runs(tmp_path):
    _, doc = _roundtrip(tmp_path)
    body = doc.paragraphs[0]
    assert _sig(body) == [("世界资源：", False, True), ("4", True, False),
                          (" 由于持续勘探。", False, False)]

    cells = [c.paragraphs[0] for c in doc.tables[0].rows[0].cells]
    # Untouched cells keep their original runs exactly
    assert _sig(cells[0]) == [("e", True, False), ("2,630", False, False)]
    assert _sig(cells[1]) == [("W", False, False)]
    assert _sig(cells[2]) == [("7", True, False), ("146,000", False, False)]
    assert _sig(cells[5]) == [("2023", False, False), ("e", True, False)]
    assert _sig(cells[6]) == [("NA", False, False)]
    # Translated cells: superscript only on the marker
    assert _sig(cells[3]) == [("Consumo, aparente", False, False), ("1", True, False)]
    assert _sig(cells[4]) == [("Dependencia neta de importaciones", False, False),
                              ("3", True, False),
                              (" como porcentaje del consumo aparente", False, False)]
    assert _sig(cells[8]) == [("Reservas", False, False)]

    link_para = doc.paragraphs[1]
    assert [(r.text, bool(r.font.superscript)) for r in link_para.runs if r.text] == \
        [("Texto", False), ("5", True)]
    assert "enlace" in link_para._p.xml and "{{" not in link_para._p.xml


def test_bilingual_keeps_markers(tmp_path):
    _, doc = _roundtrip(tmp_path, bilingual=True)
    cell = doc.tables[0].rows[0].cells[3].paragraphs[0]
    sups = [r.text for r in cell.runs if r.font.superscript]
    assert sups == ["1", "1"]
    assert "{{" not in doc.element.xml


# --- Review follow-ups: deterministic "translation" = uppercase (placeholders kept)

from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def _upper_roundtrip(tmp_path, build, bilingual=False):
    tmp_path.mkdir(parents=True, exist_ok=True)
    src = str(tmp_path / "in.docx")
    doc = Document()
    build(doc)
    doc.save(src)
    src_json = extract_word_content_to_json(src, str(tmp_path / "temp"))
    with open(src_json, encoding="utf-8") as f:
        items = json.load(f)
    dst_json = str(tmp_path / "dst.json")
    with open(dst_json, "w", encoding="utf-8") as f:
        json.dump([{"count_src": i["count_src"],
                    "translated": i["value"].upper()} for i in items],
                  f, ensure_ascii=False)
    result = write_translated_content_to_word(
        src, src_json, dst_json, str(tmp_path / "temp"), str(tmp_path / "result"),
        bilingual_mode=bilingual, src_lang="en", dst_lang="es")
    return [i["value"] for i in items], Document(result)


def _sup_texts(para):
    return [r.text for r in para.runs if r.text and (
        r.font.superscript or (r.style is not None and r.style.font.superscript))]


def test_multiletter_superscript_not_moved(tmp_path):
    # "th"/"TM" are not markers: their superscript must not jump onto "AND"
    def build(doc):
        _runs([("The 5", False, False), ("th", True, False), (" item is BrandX", False, False),
               ("TM", True, False), (" and x", False, False), ("2", True, False),
               (" grows", False, False)], doc.add_paragraph())
    _, doc = _upper_roundtrip(tmp_path, build)
    assert _sup_texts(doc.paragraphs[0]) == ["2"]
    assert doc.paragraphs[0].text == "THE 5TH ITEM IS BRANDXTM AND X2 GROWS"


def _note_ref_style(doc):
    doc.styles.add_style("FootnoteReference", 2).font.superscript = True


def test_leading_footnote_reference_style_not_inherited(tmp_path):
    def build(doc):
        _note_ref_style(doc)
        p = doc.add_paragraph()
        p._p.append(parse_xml(
            f'<w:r {nsdecls("w")}><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
            '<w:footnoteReference w:id="1"/></w:r>'))
        p.add_run(" Leading ref then plain text")
    for bilingual in (False, True):
        _, doc = _upper_roundtrip(tmp_path / str(bilingual), build, bilingual)
        p = doc.paragraphs[0]
        assert p._p.xml.count("<w:footnoteReference") == 1
        assert _sup_texts(p) == []


def test_note_reference_char_style_marker(tmp_path):
    def build(doc):
        _note_ref_style(doc)
        p = doc.add_paragraph()
        p.add_run("Styled marker")
        p.add_run("2").style = "FootnoteReference"
        p.add_run(" after styled marker text")
    for bilingual in (False, True):
        values, doc = _upper_roundtrip(tmp_path / str(bilingual), build, bilingual)
        assert values == ["Styled marker{{SUP_1}} after styled marker text"]
        assert _sup_texts(doc.paragraphs[0]) == (["2", "2"] if bilingual else ["2"])


def test_literal_placeholder_text_skips_marker_path(tmp_path):
    def build(doc):
        _runs([("Use {{SUP_1}} literally", False, False), ("1", True, False),
               (" end of text", False, False)], doc.add_paragraph())
    values, doc = _upper_roundtrip(tmp_path, build)
    assert values == ["Use {{SUP_1}} literally1 end of text"]
    assert doc.paragraphs[0].text == "USE {{SUP_1}} LITERALLY1 END OF TEXT"
    assert _sup_texts(doc.paragraphs[0]) == []
